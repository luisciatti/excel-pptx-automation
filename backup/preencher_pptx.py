import pandas as pd
from pptx import Presentation
import re
import unicodedata
import os

ppt_file = r"C:\Users\320311014\Códigos\Excel\Excel_project\Excel\pptx\202602.pptx"

# Output dinâmico
base, ext = os.path.splitext(ppt_file)
output_ppt = base + "_preenchido" + ext

# Decide o Excel pelo ano do PPT
ano = base[-6:-2]  # pega os 4 dígitos do ano no nome, ex: 2025
if ano == "2025":
    excel_file = r"C:\Users\320311014\Códigos\Excel\Relatorio2025_Com_Incidentes9.xlsx"
else:
    excel_file = r"C:\Users\320311014\Códigos\Excel\Relatorio2026_Com_Incidentes9.xlsx"

print(f"DEBUG: Usando Excel {excel_file}")
print(f"DEBUG: PPT de entrada {ppt_file}")
print(f"DEBUG: PPT de saída {output_ppt}")

# Carregar dados do Excel
df = pd.read_excel(excel_file)
df["Mês"] = pd.to_datetime(df["Mês"]).dt.date
df["Cliente"] = df["Cliente"].str.strip().str.lower()

prs = Presentation(ppt_file)

mapa_meses = {
    "janeiro": 1, "fevereiro": 2, "março": 3, "abril": 4,
    "maio": 5, "junho": 6, "julho": 7, "agosto": 8,
    "setembro": 9, "outubro": 10, "novembro": 11, "dezembro": 12
}

mapa_clientes = {
    "einstein": "hospital israelita albert einstein",
    "hospital idf": "hospital idf (instituto doutor feitosa)",
    "casas andre luiz": "centro espirita nosso lar - casas andre luiz",
    "spplus": "spdm - associacao paulista para o desenvolvimento da medicina",
    "inc": "instituto de neurocirurgia capixaba ltda inc",
    "neotin": "neotin neonatal terapia intensiva ltda",
    "proton": "proton servicos radiologicos ltda",
    "hospital espirita marilia": "hospital espirita marilia",
    "instituto atena de pesquisa clinica": "instituto atena de pesquisa clinica ltda",
    "onco jaragua": "onco jaragua",
    "univille": "univille (faculdade)",
    "iot joinville": "iot joinville",
    "grupo opty": "grupo opty",
    "hns oliveira": "hns oliveira",
    "hospital das nacoes": "hospital das nacoes",
    "hospital e maternidade santa emilia": "hospital e maternidade santa emilia",
    "anestblu": "anestblu",
    "clinica santa helena": "clinica santa helena"
}

def remover_bullet(paragraph):
    try:
        pPr = paragraph._element.get_or_add_pPr()

        # remove QUALQUER tipo de bullet
        for tag in [
            "buChar",
            "buAutoNum",
            "buBlip",
            "buNone"
        ]:
            el = pPr.find(f"{{http://schemas.openxmlformats.org/drawingml/2006/main}}{tag}")
            if el is not None:
                pPr.remove(el)

        # força explicitamente SEM bullet
        from pptx.oxml.xmlchemy import OxmlElement
        buNone = OxmlElement('a:buNone')
        pPr.append(buNone)

    except:
        pass

def normalizar(texto):
    texto = texto.strip().lower()
    texto = re.sub(r"\s+", " ", texto)
    # remove acentos
    texto = ''.join(
        c for c in unicodedata.normalize('NFD', texto)
        if unicodedata.category(c) != 'Mn'
    )
    return texto

def preencher_slide(slide, cliente, mes_texto):
    cliente_norm = normalizar(cliente)
    cliente_final = mapa_clientes.get(cliente_norm, cliente_norm)

    mes_texto = mes_texto.split()[0]
    partes = mes_texto.replace("–", "-").replace("—", "-").split("-")
    if len(partes) != 2:
        print("Formato de mês inválido:", mes_texto)
        return

    mes_nome, ano = partes
    mes_num = mapa_meses.get(mes_nome.strip().lower())
    if not mes_num:
        print("Mês inválido:", mes_nome)
        return

    mes_data = pd.to_datetime(f"{ano.strip()}-{mes_num:02d}-01").date()

    print(f"DEBUG: cliente_norm={cliente_norm}, cliente_final={cliente_final}, mes_texto={mes_texto}, mes_data={mes_data}")

    dados = df[(df["Cliente"] == cliente_final) & (df["Mês"] == mes_data)]
    if dados.empty:
        print("Sem dados para:", cliente, mes_texto)
        return

    row = dados.iloc[0]
    print(f"Preenchendo slide: {cliente_final} - {mes_data}")

    valor_usuarios = int(row.get('Usuários', row.get('Usuarios', 0)))
    valor_incidentes = int(row['Incidentes'])
    valor_indisp = int(row['Minutos Indisponibilidade'])

    for idx, shape in enumerate(slide.shapes):
        if not shape.has_text_frame:
            continue
        texto = normalizar(shape.text)

        if "% disponibilidade" in texto and idx + 1 < len(slide.shapes):
            slide.shapes[idx+1].text = f"{row['Disponibilidade (%)']}"
        elif "nro de incidentes" in texto and idx + 1 < len(slide.shapes):
            slide.shapes[idx+1].text = str(valor_incidentes)
        elif texto == "disponibilidade" and idx + 1 < len(slide.shapes):
            slide.shapes[idx+1].text = f"{row['Disponibilidade']} hrs"
        elif "indisponibilidade" in texto and idx + 1 < len(slide.shapes):
            slide.shapes[idx+1].text = f"{valor_indisp} min"
        elif texto in ["usuarios", "usuários"] and idx + 1 < len(slide.shapes):
            print(f"DEBUG: Encontrado título 'Usuários' no slide {cliente_final}. Valor Excel={valor_usuarios}")
            print(f"DEBUG: Texto atual no shape abaixo de 'Usuários' = {slide.shapes[idx+1].text}")
            slide.shapes[idx+1].text = str(valor_usuarios)
            print(f"DEBUG: Texto atualizado no shape abaixo de 'Usuários' = {slide.shapes[idx+1].text}")
        elif "armazenamento" in texto:

            encontrados = []

            shapes_lista = list(slide.shapes)

            for s2 in shapes_lista[idx+1:]:
                if s2.has_text_frame and s2.text.strip() != "":
                    encontrados.append(s2)

                if len(encontrados) == 2:
                    break
                
            # DATABASE
            if len(encontrados) >= 1:
                shape_db = encontrados[0]
                shape_db.text_frame.text = f"{row['Database (Gb)']} GB"

                for p in shape_db.text_frame.paragraphs:
                    remover_bullet(p)

            # S3
            if len(encontrados) >= 2:
                shape_s3 = encontrados[1]
                shape_s3.text_frame.text = f"{row['S3 (Gb)']} GB"

                for p in shape_s3.text_frame.paragraphs:
                    remover_bullet(p)

# Loop nos slides
for slide in prs.slides:
    cliente, mes_texto = "", ""
    for shape in slide.shapes:
        if shape.has_text_frame:
            cliente_match = re.search(r"Cliente:\s*(.*?)\s*Mês:", shape.text, re.IGNORECASE)
            mes_match = re.search(r"Mês:\s*([A-Za-z]+-\d{4})", shape.text, re.IGNORECASE)

            if cliente_match:
                cliente = cliente_match.group(1).strip().lower()
            if mes_match:
                mes_texto = mes_match.group(1).strip()

    if cliente and mes_texto:
        preencher_slide(slide, cliente, mes_texto)

prs.save(output_ppt)
print("Apresentação preenchida:", output_ppt)
