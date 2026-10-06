import pandas as pd
import re
import os
from pptx import Presentation
from pptx.util import Pt
import unicodedata


# =========================
# 📂 CONFIG
# =========================
arquivos_incidentes = [
    r"C:\Users\320311014\Códigos\Excel\Excel_project\Excel\Cloud Incidents(Registro 2025).csv",
    r"C:\Users\320311014\Códigos\Excel\Excel_project\Excel\Cloud Incidents(Registro 2026).csv"
]

ppt_file = r"C:\Users\320311014\Códigos\Excel\Excel_project\Excel\pptx\202510_final.pptx"
output_ppt = r"C:\Users\320311014\Códigos\Excel\Excel_project\Excel\pptx\202510_final2.pptx"

# =========================
# 🧠 MAPA CLIENTES
# =========================
mapa_clientes = {
    "einstein": "einstein",
    "hospital israelita albert einstein": "einstein",
    "hospital espirita marilia": "hospitalespiritamarilia",
    "clinica santa helena": "santahelena",
    "hospital santa emilia": "santaemilia",
    "iot joinville": "iot",
    "grupo opty": "opty",
    "hospital das nacoes": "nacoes",
    "idf": "idf",
    "inc": "inc",
    "neotin": "neotin",
    "proton": "proton",
    "spplus": "spplus",
    "anestblu": "anestblu",
    "casas andre luiz": "casasandreluiz",
    "instituto atena de pesquisa clinica": "institutoatena",
    "onco jaragua": "oncojaragua",
    "univille": "univille",
    "hnsoliveira": "hnsoliveira"
}

# =========================
# 🔧 FUNÇÕES
# =========================
def normalizar(texto):
    texto = str(texto).strip().lower()
    texto = ''.join(
        c for c in unicodedata.normalize('NFD', texto)
        if unicodedata.category(c) != 'Mn'
    )
    texto = re.sub(r"\s+", " ", texto)
    return texto

def mapear_cliente(x):
    x = x.strip()
    if "einstein" in x:
        return "einstein"
    for chave in mapa_clientes.keys():
        if chave in x:
            return mapa_clientes[chave]
    return x

def extrair_cliente_mes(texto):
    texto = texto.replace("\n", " ")
    cliente_match = re.search(r"Cliente:\s*(.*?)\s*M[eê]s:", texto, re.IGNORECASE)
    mes_match = re.search(r"M[eê]s:\s*([A-Za-zçÇ]+[-–]\d{4})", texto, re.IGNORECASE)
    cliente = cliente_match.group(1).strip() if cliente_match else None
    mes = mes_match.group(1).replace("–", "-").strip() if mes_match else None
    return cliente, mes

# Tradução de meses
traducoes_meses = {
    "janeiro": "january",
    "fevereiro": "february",
    "marco": "march",
    "abril": "april",
    "maio": "may",
    "junho": "june",
    "julho": "july",
    "agosto": "august",
    "setembro": "september",
    "outubro": "october",
    "novembro": "november",
    "dezembro": "december"
}

# =========================
# 📊 LER INCIDENTES
# =========================
df_list = []
for caminho in arquivos_incidentes:
    if not os.path.exists(caminho):
        print(f"⚠️ Arquivo não encontrado: {caminho}")
        continue

    print(f"📥 Lendo arquivo: {caminho}")
    df = pd.read_csv(caminho, encoding="latin1")
    if "Total Time (Minutes)" in df.columns:
        df["Total Time (Minutes)"] = pd.to_numeric(df["Total Time (Minutes)"], errors="coerce").fillna(0)

    df["Customer_norm"] = df["Customer"].apply(normalizar).apply(mapear_cliente)
    df["Year"] = pd.to_numeric(df["Year"], errors="coerce").astype("Int64")  # ✅ seguro
    df["Mes_norm"] = df["Month"].apply(normalizar) + "-" + df["Year"].astype(str)


    df_list.append(df)

df_all = pd.concat(df_list, ignore_index=True)
print(f"✅ Total de incidentes carregados: {len(df_all)}")

# 🔎 Debug extra: mostrar alguns clientes
for cliente in ["einstein","iot","santahelena","opty"]:
    print(f"\n📊 Amostra de {cliente} no CSV (já normalizado):")
    print(df_all[df_all["Customer_norm"] == cliente][["Incident SNOW","Date","Month","Year","Customer","Mes_norm"]].head(10))

# =========================
# 📂 PPT
# =========================
prs = Presentation(ppt_file)

# =========================
# 🔄 LOOP SLIDES
# =========================
for slide in prs.slides:
    texto_slide = " ".join([shape.text for shape in slide.shapes if shape.has_text_frame])
    cliente, mes = extrair_cliente_mes(texto_slide)

    if not cliente or not mes:
        continue

    cliente_norm = normalizar(cliente)
    cliente_norm = mapa_clientes.get(cliente_norm, cliente_norm)
    mes_norm = normalizar(mes)

    # Flexibilizar meses (português/inglês)
    possiveis_meses = [mes_norm]
    for pt, en in traducoes_meses.items():
        if pt in mes_norm:
            possiveis_meses.append(mes_norm.replace(pt, en))

    print(f"\n🔎 Slide: Cliente={cliente} | Normalizado={cliente_norm} | Mês={mes} | Possíveis={possiveis_meses}")

    filtro = df_all[(df_all["Customer_norm"] == cliente_norm) & (df_all["Mes_norm"].isin(possiveis_meses))]

    print(f"   ➡️ Incidentes encontrados: {len(filtro)}")

    for shape in slide.shapes:
        if not shape.has_table:
            continue

        tabela = shape.table

        # Limpar conteúdo das linhas (menos cabeçalho)
        for r in range(1, len(tabela.rows)):
            for c in range(len(tabela.columns)):
                tabela.rows[r].cells[c].text = ""

        if filtro.empty:
            print("   ⚠️ Nenhum incidente encontrado, escrevendo 'Sem registros'")
            if len(tabela.rows) > 1:
                row = tabela.rows[1]
                row.cells[0].text = "Resumo"
                row.cells[5].text = "Sem registros de incidentes no mês."
                for c in row.cells:
                    for paragraph in c.text_frame.paragraphs:
                        for run in paragraph.runs:
                             run.font.size = Pt(10)  # ou 9, ou 8, conforme desejar
        else:
            i = 1
            for _, inc in filtro.iterrows():
                if i >= len(tabela.rows):
                    break
                print(f"   ✅ Escrevendo incidente: {inc['Incident SNOW']} - {inc['Problem Description']}")
                row = tabela.rows[i]
                row.cells[0].text = str(inc["Date"])
                row.cells[1].text = str(inc.get("Recurso", ""))
                row.cells[2].text = str(inc["Starting Date/Time"])
                row.cells[3].text = str(inc["Ending Date/Time"])
                row.cells[4].text = str(int(inc["Total Time (Minutes)"])) + " min"
                row.cells[5].text = str(inc["Problem Description"])
                row.cells[6].text = str(inc["Correction Action"])
                i += 1
                for c in row.cells:
                    for paragraph in c.text_frame.paragraphs:
                        for run in paragraph.runs:
                             run.font.size = Pt(10)  # ou 9, ou 8, conforme desejar


# =========================
# 💾 SALVAR
# =========================
prs.save(output_ppt)
print("\n🚀 PPT FINAL GERADO:")
print(output_ppt)
