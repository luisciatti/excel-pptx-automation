import pandas as pd
import os
import re
from collections import Counter
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN
import unicodedata

# =========================
# 📂 CONFIG
# =========================
pasta_logs = r"C:\Users\320311014\Códigos\Excel\Excel_project\Excel\logs"
ppt_file = r"C:\Users\320311014\Códigos\Excel\Excel_project\Excel\pptx\202602_preenchido.pptx"
output_ppt = r"C:\Users\320311014\Códigos\Excel\Excel_project\Excel\pptx\202602_final.pptx"

# =========================
# 🧠 MAPA MESES
# =========================
mapa_meses = {
    "01": "janeiro", "02": "fevereiro", "03": "março", "04": "abril",
    "05": "maio", "06": "junho", "07": "julho", "08": "agosto",
    "09": "setembro", "10": "outubro", "11": "novembro", "12": "dezembro"
}

# =========================
# 🧠 MAPA CLIENTES
# =========================
mapa_clientes = {
    "hospital e maternidade santa emilia": "santaemilia",
    "santa emilia": "santaemilia",
    "clinica santa helena": "santahelena",
    "iot joinville": "iot" ,
    "grupo opty": "opty",
    "hns oliveira": "hnsoliveira",
    "hospital das nacoes": "nacoes",
    "hospital idf": "idf",
    "inc": "inc",
    "neotin": "neotin",
    "proton": "proton",
    "spplus": "spplus",
    "anestblu": "anestblu",
    "casas andre luiz": "casasandreluiz",
    "hospital espirita marilia": "hospitalespiritamarilia",
    "instituto atena de pesquisa clinica": "institutoatena",
    "onco jaragua": "oncojaragua",
    "univille": "univille",
    "cmabc": "cmabc",
    "einstein": "einstein"
}

# =========================
# 🔧 FUNÇÕES
# =========================
def corrigir_texto(texto):
    try:
        return texto.encode("latin1").decode("utf-8")
    except:
        return texto
    
def normalizar(texto):
    texto = str(texto).strip().lower()
    texto = ''.join(
        c for c in unicodedata.normalize('NFD', texto)
        if unicodedata.category(c) != 'Mn'
    )
    texto = re.sub(r"\s+", " ", texto)
    return texto

def extrair_cliente_mes(texto):
    texto = texto.replace("\n", " ")

    cliente_match = re.search(r"Cliente:\s*(.*?)\s*M[eê]s:", texto, re.IGNORECASE)
    mes_match = re.search(r"M[eê]s:\s*([A-Za-zçÇ]+[-–]\d{4})", texto, re.IGNORECASE)

    cliente = cliente_match.group(1).strip() if cliente_match else None
    mes = mes_match.group(1).replace("–", "-").strip() if mes_match else None

    return cliente, mes

def remover_bullet(paragraph):
    try:
        pPr = paragraph._element.get_or_add_pPr()
        bu = pPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar")
        if bu is not None:
            pPr.remove(bu)
    except:
        pass

# =========================
# 📊 LER LOGS
# =========================
dados = []

for arquivo in os.listdir(pasta_logs):

    if not arquivo.endswith(".csv"):
        continue

    match = re.search(r"(.*?)-throttling-log-(\d{4})-(\d{2})", arquivo)

    if not match:
        continue

    cliente = normalizar(match.group(1))
    cliente = mapa_clientes.get(cliente, cliente)

    ano = match.group(2)
    mes_num = match.group(3)
    mes_nome = mapa_meses.get(mes_num, mes_num)

    caminho = os.path.join(pasta_logs, arquivo)

    try:
        df = pd.read_csv(caminho, encoding="latin1")
    except:
        continue

    if "QTD" not in df.columns or "FUNCAO" not in df.columns:
        continue

    df["QTD"] = pd.to_numeric(df["QTD"], errors="coerce").fillna(0)

    total_qtd = int(df["QTD"].sum())
    funcoes = df["FUNCAO"].dropna().astype(str).apply(corrigir_texto).tolist()
    top_funcoes = [corrigir_texto(f).strip().capitalize() for f, _ in Counter(funcoes).most_common(5)]

    dados.append({
        "cliente": cliente,
        "mes": normalizar(f"{mes_nome}-{ano}"),
        "total_qtd": total_qtd,
        "top5": top_funcoes
    })

df_final = pd.DataFrame(dados)

print("\n📊 BASE FINAL:")
print(df_final)

# =========================
# 📂 PPT
# =========================
prs = Presentation(ppt_file)

# =========================
# 🔄 LOOP SLIDES
# =========================
for slide in prs.slides:

    texto_slide = " ".join(
        [shape.text for shape in slide.shapes if shape.has_text_frame]
    )

    cliente, mes = extrair_cliente_mes(texto_slide)

    if not cliente or not mes:
        continue

    cliente_norm = normalizar(cliente)
    cliente_norm = mapa_clientes.get(cliente_norm, cliente_norm)

    mes_norm = normalizar(mes)

    filtro = df_final[
        (df_final["cliente"] == cliente_norm) &
        (df_final["mes"] == mes_norm)
    ]

    if filtro.empty:
        print(f"⚠️ Sem dados: {cliente_norm} - {mes_norm}")
        continue

    row = filtro.iloc[0]

    total_qtd = row["total_qtd"]
    top_funcoes = row["top5"]

    print(f"\n✅ Preenchendo: {cliente_norm} - {mes_norm}")

    for shape in slide.shapes:

        if not shape.has_text_frame:
            continue

        texto = shape.text

        # =========================
        # 🔥 QTD (SÓ O CERTO)
        # =========================
        if "Ações Pró-ativas no Ambiente" in texto:

            ref_shape = shape
            ref_top = ref_shape.top

            candidato = None
            menor_dist = None

            for s2 in slide.shapes:

                if not s2.has_text_frame or s2 == ref_shape:
                    continue

                txt = s2.text.strip()

                if not re.fullmatch(r"\d+", txt):
                    continue

                dist = abs(s2.top - ref_top)

                if menor_dist is None or dist < menor_dist:
                    menor_dist = dist
                    candidato = s2

            if candidato:
                candidato.text_frame.text = str(total_qtd)

                for p in candidato.text_frame.paragraphs:
                    p.font.size = Pt(10)
                    p.font.bold = False
                    p.alignment = PP_ALIGN.LEFT

        # =========================
        # 🔥 TOP 5
        # =========================
        if "Top 5 - sessões encerradas" in texto:
        
           tf = shape.text_frame
        
           # limpa de forma segura
           tf.clear()
        
           # título
           p_title = tf.paragraphs[0]
           p_title.text = "Top 5 - sessões encerradas"
           p_title.font.size = Pt(12)
           p_title.font.bold = True
           remover_bullet(p_title)
        
           # =========================
           # SEM DADOS
           # =========================
           if len(top_funcoes) == 0:
               p = tf.add_paragraph()
               p.text = ""
               p.font.size = Pt(11)
               p.alignment = PP_ALIGN.CENTER
               remover_bullet(p)
        
           # =========================
           # COM DADOS
           # =========================
           else:
               for f in top_funcoes:
                   p = tf.add_paragraph()
                   p.text = ""
                   p.text = f
                   p.font.size = Pt(11)
                   p.alignment = PP_ALIGN.LEFT
                   remover_bullet(p)
        
           
# =========================
# 💾 SALVAR
# =========================
prs.save(output_ppt)

print("\n🚀 PPT FINAL GERADO:")
print(output_ppt)